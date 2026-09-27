#!/usr/bin/env python3
"""Build a review workbook only. No account access, joining, posting or scheduling."""
import csv,json
from pathlib import Path
from urllib.parse import quote
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.datavalidation import DataValidation

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'deliverables';OUT.mkdir(exist_ok=True)
articles=json.loads((ROOT/'scripts/section-launch-articles.json').read_text())
config={
'los-angeles':('Los Angeles','City services; MyLA311','Where should an LA neighborhood problem actually be reported?','Our CaliReporter guide explains the MyLA311 starting point for city service requests. What part of reporting a neighborhood problem is hardest to navigate?',['Los Angeles neighborhood city services','Los Angeles community maintenance','Los Angeles MyLA311 residents']),
'san-diego':('San Diego','City services; Get It Done','Pothole, graffiti or missed pickup: where does your report go?','CaliReporter put together a guide to San Diego’s Get It Done service. Which type of routine neighborhood issue is hardest to report clearly?',['San Diego neighborhood city services','San Diego Get It Done community','San Diego neighborhood maintenance']),
'san-jose':('San Jose','City services; San Jose 311','Do you know how to track a San Jose service request?','Our CaliReporter guide covers reporting and tracking routine issues through San Jose 311. What would make the reporting process easier to understand?',['San Jose neighborhood 311','San Jose community city services','San Jose neighborhood maintenance']),
'san-francisco':('San Francisco','City information; 311','Not sure which San Francisco department handles your question?','CaliReporter explains how 311 can help residents find the right city service. Which city process would you like to see explained next?',['San Francisco neighborhood 311','San Francisco community city services','San Francisco neighborhood information']),
'fresno':('Fresno','City services; FresGO','Do you use FresGO, call 311, or start somewhere else?','Our CaliReporter guide outlines Fresno’s options for reporting routine neighborhood concerns. Which reporting channel do you find easiest to use?',['Fresno neighborhood FresGO','Fresno community city services','Fresno neighborhood maintenance']),
'health':('California / statewide','Nutrition literacy; food labels','Are you adding the two sugar numbers on a food label together?','CaliReporter’s FDA-sourced explainer covers total sugars versus added sugars. Which part of a Nutrition Facts label would you like explained next?',['California nutrition education community','California cooking food labels','California healthy eating discussion']),
'relationships':('California / statewide','Communication; boundaries','What makes a difficult conversation more respectful?','Our CaliReporter guide looks at NIH advice on clear requests, listening and boundaries. What communication habit do you find most helpful?',['California relationships communication discussion','California community communication boundaries','California social wellness discussion']),
}
wb=Workbook();readme=wb.active;readme.title='Start Here'
readme.append(['Item','Instructions / status'])
for row in [
('Status','DRAFT WORKBOOK ONLY. No new agent, scheduled search, account joining or posting has been enabled.'),
('Purpose','Prepare clearly attributed CaliReporter posts and find relevant groups by the actual article topic and place.'),
('Article Drafts','Seven editable starter posts using already published, source-linked explainers. Edit Draft headline and Draft post text.'),
('Group Searches','Three search terms per article with clickable Facebook group-search URLs. Search links are not verified group results.'),
('Group Candidates','Two public-source discovery leads. Current membership eligibility and publisher-link rules remain unverified. No approved destinations.'),
('Proposed workflow','New article -> extract location/topic -> discover groups -> verify fit, membership eligibility and group rules -> prepare spreadsheet draft -> distribute only through identified publisher/owner accounts to eligible groups.'),
('Account routing','Two profile slots are unassigned. Identify genuine owner/publisher accounts and their visible CaliReporter affiliation before connecting any automation.'),
('Group choice','Up to three eligible groups for the primary account, subject to group rules and frequency limits. Discovery is based on relevance, not a requirement to fill three slots.'),
('Duplicate handling','One article/group combination across all participating accounts; no duplicate promotion through different profiles.'),
('Joining','Check membership criteria, answer questions truthfully, and do not claim local residency or professional credentials that the account holder does not have.'),
('Copy standards','Accurate headline and open question; disclose CaliReporter. Do not invent facts, controversy, quotations or independent local endorsements.'),
('Posting Log','Empty until real actions occur. Save actual post URL/result, including pending moderation or failed submission.'),
('Research date','2026-09-27. Public community websites supplied the two leads. Facebook rules could not be verified from public browsing.'),
]:readme.append(row)
ws=wb.create_sheet('Article Drafts');headers=['Article ID','Article URL','Published title','Location','Topic','Draft headline','Draft post text','Source URL','Assigned profile','Chosen group URL','Rules checked','Draft status','Posted URL','Notes'];ws.append(headers)
for i,a in enumerate(articles,1):
 place,topic,headline,body,queries=config[a['category']];url='https://calireporter.com/article/'+a['slug']
 ws.append([f'CR-{i:03}',url,a['title'],place,topic,headline,headline+'\n\n'+body+'\n\n'+url,a['source_url'],'','','No','Draft','','Starter copy; review current article before use.'])
search=wb.create_sheet('Group Searches');search.append(['Article ID','Location','Topic','Search phrase','Facebook group search URL','Status','Candidate group URL','Relevance notes'])
for i,a in enumerate(articles,1):
 place,topic,headline,body,queries=config[a['category']]
 for q in queries:search.append([f'CR-{i:03}',place,topic,q,'https://www.facebook.com/search/groups/?q='+quote(q),'Research queued','','Verify locality and topic; publisher links must be allowed.'])
groups=wb.create_sheet('Group Candidates');groups.append(['Candidate ID','Group name','Group URL','Place','Topic','Discovery source','Match assessment','Current group rules','Membership eligibility','Review status','Profile','Notes'])
groups.append(['G-001','Los Angeles Neighborhood Council Sustainability Alliance','https://www.facebook.com/groups/LosAngelesNCSA','Los Angeles','Sustainability; neighborhood issues','https://www.ncsa.la/about','Potential for environmental/neighborhood-service stories; assess each article','Unknown','Unknown','Needs rules review','','Community website links this group. Broad city-service guide may be too general; do not assume fit or permission.'])
groups.append(['G-002','Fans of Japantown San Jose','https://www.facebook.com/groups/222379767772678/','San Jose / Japantown','Neighborhood information; culture','https://jtown.org/tanoshii','Potential for directly relevant Japantown service/culture articles','Unknown','Unknown','Needs rules review','','Jtown.org links this group. Citywide 311 guide needs a clear neighborhood connection and rule review. Facebook research fetch did not expose current rules.'])
accounts=wb.create_sheet('Account Setup');accounts.append(['Slot','Intended role','AdsPower profile ID','Account display name','CaliReporter affiliation confirmed','Allowed group URLs','Status'])
accounts.append(['Primary','Distribute to up to 3 eligible groups per article','','','No','','Waiting for identity'])
accounts.append(['Discovery','Find topic/place matches; distribute only after eligibility/rules checks','','','No','','Waiting for identity'])
log=wb.create_sheet('Posting Log');log.append(['Article ID','Group URL','Profile ID','Attempted at','Outcome','Facebook post URL','Notes'])
for sheet in wb:
 sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
 for c in sheet[1]:c.fill=PatternFill('solid',fgColor='1E293B');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True,vertical='center')
 sheet.row_dimensions[1].height=34
 for row in sheet.iter_rows(min_row=2):
  for c in row:
   c.alignment=Alignment(wrap_text=True,vertical='top')
   if isinstance(c.value,str) and c.value.startswith('https://'):
    c.hyperlink=c.value;c.font=Font(color='1D4ED8',underline='single')
  sheet.row_dimensions[row[0].row].height=90 if sheet.title!='Article Drafts' else 180
 for col in sheet.columns:
  name=col[0].value or '';sheet.column_dimensions[col[0].column_letter].width=52 if any(x in name for x in ['text','Instructions','Notes','URL','assessment','source']) else 28
readme.column_dimensions['A'].width=25;readme.column_dimensions['B'].width=110
ws.column_dimensions['G'].width=85;ws.column_dimensions['F'].width=55
for column,choices in [('K','No,Yes,Not allowed'),('L','Draft,Ready for review,Approved,Rejected,Posted')]:
 dv=DataValidation(type='list',formula1='"'+choices+'"');dv.errorTitle='Choose a listed status';dv.error='Use the dropdown status.';dv.showErrorMessage=True;dv.errorStyle='stop';ws.add_data_validation(dv);dv.add(f'{column}2:{column}10000')
filename=OUT/'CaliReporter-Facebook-Planning.xlsx';wb.save(filename)
with (OUT/'CaliReporter-Facebook-Drafts.csv').open('w',newline='',encoding='utf-8-sig') as f:
 csv.writer(f).writerows(ws.values)
check=load_workbook(filename);assert check['Article Drafts'].max_row==8;assert check['Group Searches'].max_row==22;assert check['Group Candidates'].max_row==3;assert check['Posting Log'].max_row==1
assert all(check['Article Drafts'].cell(i,12).value=='Draft' for i in range(2,9))
assert all('CaliReporter' in check['Article Drafts'].cell(i,7).value for i in range(2,9))
print(json.dumps({'file':str(filename),'sheets':check.sheetnames,'drafts':7,'search_queries':21,'unverified_candidates':2,'posts_sent':0}))
